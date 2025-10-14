from flask import Flask, render_template, request, send_file, url_for
from PIL import Image
import io
import os
import uuid
import tempfile
import fitz
import zipfile

app = Flask(__name__, template_folder='../templates', static_folder='../static')

file_storage = {}

def jpg_to_png(image_file, original_filename):
    img = Image.open(image_file)
    img_io = io.BytesIO()
    img.save(img_io, format='PNG')
    img_io.seek(0)
    base_name = os.path.splitext(original_filename)[0]
    return img_io, f'{base_name}.png'

def png_to_jpg(image_file, original_filename):
    img = Image.open(image_file)
    img_io = io.BytesIO()
    img.convert('RGB').save(img_io, format='JPEG')
    img_io.seek(0)
    base_name = os.path.splitext(original_filename)[0]
    return img_io, f'{base_name}.jpg'

def jpg_to_pdf(image_file, original_filename):
    img = Image.open(image_file)
    pdf_io = io.BytesIO()
    img.convert('RGB').save(pdf_io, format='PDF')
    pdf_io.seek(0)
    base_name = os.path.splitext(original_filename)[0]
    return pdf_io, f'{base_name}.pdf'

def png_to_pdf(image_file, original_filename):
    img = Image.open(image_file)
    pdf_io = io.BytesIO()
    img.convert('RGB').save(pdf_io, format='PDF')
    pdf_io.seek(0)
    base_name = os.path.splitext(original_filename)[0]
    return pdf_io, f'{base_name}.pdf'

def word_to_pdf(word_file, original_filename):
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, ListFlowable, ListItem
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch, pt
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    temp_word_path = None
    temp_pdf_path = None
    try:
        temp_word_path = tempfile.mktemp(suffix='.docx')
        word_file.save(temp_word_path)
        
        doc = Document(temp_word_path)
        temp_pdf_path = tempfile.mktemp(suffix='.pdf')
        
        sections = doc.sections[0]
        page_width = sections.page_width.inches * inch
        page_height = sections.page_height.inches * inch
        left_margin = sections.left_margin.inches * inch if sections.left_margin else 72
        right_margin = sections.right_margin.inches * inch if sections.right_margin else 72
        top_margin = sections.top_margin.inches * inch if sections.top_margin else 72
        bottom_margin = sections.bottom_margin.inches * inch if sections.bottom_margin else 72
        
        pdf_doc = SimpleDocTemplate(
            temp_pdf_path, 
            pagesize=(page_width, page_height),
            rightMargin=right_margin,
            leftMargin=left_margin,
            topMargin=top_margin,
            bottomMargin=bottom_margin
        )
        
        styles = getSampleStyleSheet()
        story = []
        
        for paragraph in doc.paragraphs:
            text = paragraph.text
            
            if not text.strip():
                spacing = 6
                if paragraph.paragraph_format.space_after:
                    spacing = paragraph.paragraph_format.space_after.pt
                elif paragraph.paragraph_format.space_before:
                    spacing = paragraph.paragraph_format.space_before.pt
                story.append(Spacer(1, spacing))
                continue
            
            font_size = 11
            font_name = 'Helvetica'
            is_bold = False
            is_italic = False
            
            if paragraph.runs:
                first_run = paragraph.runs[0]
                if first_run.font.size:
                    font_size = first_run.font.size.pt
                if first_run.bold:
                    is_bold = True
                if first_run.italic:
                    is_italic = True
            
            if is_bold and is_italic:
                font_name = 'Helvetica-BoldOblique'
            elif is_bold:
                font_name = 'Helvetica-Bold'
            elif is_italic:
                font_name = 'Helvetica-Oblique'
            
            alignment = TA_LEFT
            if paragraph.alignment == 1:
                alignment = TA_CENTER
            elif paragraph.alignment == 2:
                alignment = TA_RIGHT
            elif paragraph.alignment == 3:
                alignment = TA_JUSTIFY
            
            line_spacing = 1.15
            if paragraph.paragraph_format.line_spacing:
                if isinstance(paragraph.paragraph_format.line_spacing, float):
                    line_spacing = paragraph.paragraph_format.line_spacing
            
            space_before = 0
            space_after = 6
            if paragraph.paragraph_format.space_before:
                space_before = paragraph.paragraph_format.space_before.pt
            if paragraph.paragraph_format.space_after:
                space_after = paragraph.paragraph_format.space_after.pt
            
            left_indent = 0
            first_line_indent = 0
            if paragraph.paragraph_format.left_indent:
                left_indent = paragraph.paragraph_format.left_indent.pt
            if paragraph.paragraph_format.first_line_indent:
                first_line_indent = paragraph.paragraph_format.first_line_indent.pt
            
            style_name = paragraph.style.name
            if style_name.startswith('Heading'):
                if '1' in style_name:
                    base_style = styles['Heading1']
                elif '2' in style_name:
                    base_style = styles['Heading2']
                elif '3' in style_name:
                    base_style = styles['Heading3']
                else:
                    base_style = styles['Heading4']
            elif style_name == 'List Bullet' or style_name == 'List Number':
                base_style = styles['Normal']
            else:
                base_style = styles['Normal']
            
            custom_style = ParagraphStyle(
                f'custom_{id(paragraph)}',
                parent=base_style,
                fontSize=font_size,
                fontName=font_name,
                alignment=alignment,
                leading=font_size * line_spacing,
                spaceBefore=space_before,
                spaceAfter=space_after,
                leftIndent=left_indent,
                firstLineIndent=first_line_indent
            )
            
            formatted_text = text
            for run in paragraph.runs:
                run_text = run.text
                if run.bold and run.italic:
                    formatted_text = formatted_text.replace(run_text, f'<b><i>{run_text}</i></b>', 1)
                elif run.bold:
                    formatted_text = formatted_text.replace(run_text, f'<b>{run_text}</b>', 1)
                elif run.italic:
                    formatted_text = formatted_text.replace(run_text, f'<i>{run_text}</i>', 1)
                elif run.underline:
                    formatted_text = formatted_text.replace(run_text, f'<u>{run_text}</u>', 1)
            
            para = Paragraph(formatted_text, custom_style)
            story.append(para)
        
        for table in doc.tables:
            data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                data.append(row_data)
            
            if data:
                col_widths = None
                if table.rows:
                    num_cols = len(table.rows[0].cells)
                    available_width = page_width - left_margin - right_margin
                    col_widths = [available_width / num_cols] * num_cols
                
                t = Table(data, colWidths=col_widths)
                t.setStyle(TableStyle([
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ]))
                story.append(t)
                story.append(Spacer(1, 12))
        
        pdf_doc.build(story)
        
        with open(temp_pdf_path, 'rb') as f:
            pdf_io = io.BytesIO(f.read())
        
        base_name = os.path.splitext(original_filename)[0]
        return pdf_io, f'{base_name}.pdf'
    except Exception as e:
        print(f"Error converting Word to PDF: {e}")
        raise
    finally:
        if temp_word_path and os.path.exists(temp_word_path):
            try:
                os.remove(temp_word_path)
            except:
                pass
        if temp_pdf_path and os.path.exists(temp_pdf_path):
            try:
                os.remove(temp_pdf_path)
            except:
                pass

def pdf_to_word(pdf_file, original_filename):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor, Cm
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
    
    temp_pdf_path = None
    temp_word_path = None
    try:
        temp_pdf_path = tempfile.mktemp(suffix='.pdf')
        with open(temp_pdf_path, 'wb') as f:
            f.write(pdf_file.read())
        
        pdf_document = fitz.open(temp_pdf_path)
        doc = Document()
        
        page = pdf_document[0]
        page_rect = page.rect
        section = doc.sections[0]
        section.page_width = Inches(page_rect.width / 72)
        section.page_height = Inches(page_rect.height / 72)
        
        for page_num in range(len(pdf_document)):
            page = pdf_document.load_page(page_num)
            blocks = page.get_text("dict")["blocks"]
            
            for block in blocks:
                if block.get("type") == 0:
                    block_bbox = block.get("bbox", [0, 0, 0, 0])
                    
                    for line in block.get("lines", []):
                        line_text = ""
                        spans = line.get("spans", [])
                        
                        if not spans:
                            continue
                        
                        para = doc.add_paragraph()
                        
                        for span in spans:
                            span_text = span.get("text", "")
                            if not span_text:
                                continue
                            
                            run = para.add_run(span_text)
                            
                            font_size = span.get("size", 11)
                            font_name = span.get("font", "Calibri")
                            font_flags = span.get("flags", 0)
                            color = span.get("color", 0)
                            
                            run.font.size = Pt(font_size)
                            
                            try:
                                if "Bold" in font_name or font_flags & 2**4:
                                    run.font.bold = True
                                if "Italic" in font_name or "Oblique" in font_name or font_flags & 2**1:
                                    run.font.italic = True
                                
                                clean_font_name = font_name.split('-')[0] if '-' in font_name else font_name
                                clean_font_name = clean_font_name.replace('Bold', '').replace('Italic', '').replace('Oblique', '').strip()
                                if clean_font_name:
                                    run.font.name = clean_font_name
                            except:
                                pass
                            
                            if color != 0:
                                try:
                                    r = (color >> 16) & 0xFF
                                    g = (color >> 8) & 0xFF
                                    b = color & 0xFF
                                    run.font.color.rgb = RGBColor(r, g, b)
                                except:
                                    pass
                        
                        if para.text.strip():
                            line_bbox = line.get("bbox", [0, 0, 0, 0])
                            x_offset = line_bbox[0]
                            
                            if x_offset > 100:
                                para.paragraph_format.left_indent = Pt(x_offset - 72)
                            
                            line_height = line_bbox[3] - line_bbox[1]
                            if spans:
                                avg_font_size = sum(s.get("size", 11) for s in spans) / len(spans)
                                spacing_ratio = line_height / avg_font_size if avg_font_size > 0 else 1.0
                                
                                if 0.8 <= spacing_ratio <= 2.5:
                                    para.paragraph_format.line_spacing = spacing_ratio
                        else:
                            para._element.getparent().remove(para._element)
                
                elif block.get("type") == 1:
                    try:
                        img_dict = block
                        xref = img_dict.get("image")
                        if xref:
                            base_image = pdf_document.extract_image(xref)
                            image_bytes = base_image["image"]
                            
                            temp_img = tempfile.mktemp(suffix=f'.{base_image["ext"]}')
                            with open(temp_img, 'wb') as img_file:
                                img_file.write(image_bytes)
                            
                            bbox = block.get("bbox", [0, 0, 0, 0])
                            img_width = bbox[2] - bbox[0]
                            img_height = bbox[3] - bbox[1]
                            
                            width_inches = img_width / 72
                            height_inches = img_height / 72
                            
                            max_width = section.page_width.inches - 2
                            if width_inches > max_width:
                                ratio = max_width / width_inches
                                width_inches = max_width
                                height_inches = height_inches * ratio
                            
                            doc.add_picture(temp_img, width=Inches(width_inches))
                            
                            try:
                                os.remove(temp_img)
                            except:
                                pass
                    except Exception as e:
                        print(f"Error adding image: {e}")
                        pass
            
            if page_num < len(pdf_document) - 1:
                doc.add_page_break()
        
        pdf_document.close()
        
        temp_word_path = tempfile.mktemp(suffix='.docx')
        doc.save(temp_word_path)
        
        with open(temp_word_path, 'rb') as f:
            docx_io = io.BytesIO(f.read())
        
        base_name = os.path.splitext(original_filename)[0]
        return docx_io, f'{base_name}.docx'
    except Exception as e:
        print(f"Error converting PDF to Word: {e}")
        raise
    finally:
        if temp_pdf_path and os.path.exists(temp_pdf_path):
            try:
                os.remove(temp_pdf_path)
            except:
                pass
        if temp_word_path and os.path.exists(temp_word_path):
            try:
                os.remove(temp_word_path)
            except:
                pass

def pdf_to_jpg(pdf_file, original_filename):
    pdf_document = fitz.open(stream=pdf_file.read(), filetype="pdf")
    zip_io = io.BytesIO()
    try:
        with zipfile.ZipFile(zip_io, mode='w', compression=zipfile.ZIP_DEFLATED) as zip_file:
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                pix = page.get_pixmap(dpi=150)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                img_io = io.BytesIO()
                img.save(img_io, format='JPEG', quality=85)
                img_io.seek(0)
                zip_file.writestr(f'{os.path.splitext(original_filename)[0]}_page_{page_num + 1}.jpg', img_io.read())
        zip_io.seek(0)
        return zip_io, f'{os.path.splitext(original_filename)[0]}.zip'
    finally:
        pdf_document.close()

def pdf_to_png(pdf_file, original_filename):
    pdf_document = fitz.open(stream=pdf_file.read(), filetype="pdf")
    zip_io = io.BytesIO()
    try:
        with zipfile.ZipFile(zip_io, mode='w', compression=zipfile.ZIP_DEFLATED) as zip_file:
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                pix = page.get_pixmap(dpi=150)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                img_io = io.BytesIO()
                img.save(img_io, format='PNG', optimize=True)
                img_io.seek(0)
                zip_file.writestr(f'{os.path.splitext(original_filename)[0]}_page_{page_num + 1}.png', img_io.read())
        zip_io.seek(0)
        return zip_io, f'{os.path.splitext(original_filename)[0]}.zip'
    finally:
        pdf_document.close()

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        try:
            file = request.files.get('file')
            if not file or not file.filename:
                return render_template('index.html', show_button=False, error='Please select a file')
            
            file_type = request.form.get('file_type')
            if not file_type:
                return render_template('index.html', show_button=False, error='Please select a conversion type')
            
            original_filename = file.filename
            file_id = str(uuid.uuid4())
            
            if file_type == 'jpg_to_png':
                img_io, filename = jpg_to_png(file, original_filename)
                file_storage[file_id] = (img_io, 'image/png', filename)

            elif file_type == 'jpg_to_pdf':
                pdf_io, filename = jpg_to_pdf(file, original_filename)
                file_storage[file_id] = (pdf_io, 'application/pdf', filename)

            elif file_type == 'png_to_jpg':
                img_io, filename = png_to_jpg(file, original_filename)
                file_storage[file_id] = (img_io, 'image/jpeg', filename)

            elif file_type == 'png_to_pdf':
                pdf_io, filename = png_to_pdf(file, original_filename)
                file_storage[file_id] = (pdf_io, 'application/pdf', filename)

            elif file_type == 'pdf_to_jpg':
                zip_io, filename = pdf_to_jpg(file, original_filename)
                file_storage[file_id] = (zip_io, 'application/zip', filename)

            elif file_type == 'pdf_to_png':
                zip_io, filename = pdf_to_png(file, original_filename)
                file_storage[file_id] = (zip_io, 'application/zip', filename)
            
            else:
                return render_template('index.html', show_button=False, error='Invalid conversion type')

            return render_template('index.html', download_url=url_for('download', file_id=file_id), show_button=True)
        
        except Exception as e:
            print(f"Error during conversion: {str(e)}")
            return render_template('index.html', show_button=False, error=f'Conversion failed: {str(e)}')
    
    return render_template('index.html', show_button=False)

@app.route('/download/<file_id>')
def download(file_id):
    file_data = file_storage.get(file_id)
    if file_data:
        file_io, mime_type, filename = file_data
        file_io.seek(0)
        response = send_file(file_io, mimetype=mime_type, as_attachment=True, download_name=filename)
        response.headers["Refresh"] = "1; url='/'"
        del file_storage[file_id]
        return response
    return 'File not found', 404

# For Vercel
app = app
